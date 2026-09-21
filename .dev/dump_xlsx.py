import openpyxl,sys,glob
from openpyxl.utils import get_column_letter as L
sys.stdout.reconfigure(encoding='utf-8')
def dump(path,sheet=0,maxcol=26,skip_hidden=True):
    wb=openpyxl.load_workbook(path,data_only=True); ws=wb.worksheets[sheet]
    hc={k for k,v in ws.column_dimensions.items() if v.hidden}
    for r in range(1,ws.max_row+1):
        if skip_hidden and ws.row_dimensions[r].hidden: continue
        out=[]
        for c in range(1,maxcol+1):
            col=L(c)
            if col in hc: continue
            v=ws.cell(r,c).value
            if v is None or (isinstance(v,str) and not v.strip()): continue
            if isinstance(v,float): v=round(v,3)
            out.append(f'{col}={v}')
        if out: print(r,' | '.join(out))
if __name__=='__main__':
    dump(glob.glob(sys.argv[1])[0],int(sys.argv[2]) if len(sys.argv)>2 else 0)
