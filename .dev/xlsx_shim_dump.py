# -*- coding: utf-8 -*-
# ตัวช่วยของ .dev/xlsx.full.min.js (ดู docstring ในไฟล์นั้น) — อ่านไฟล์ .xlsx ด้วย openpyxl แล้วทิ้งเป็น JSON
# รูปร่างที่พอให้ shim ฝั่ง Node ประกอบเป็นอ็อบเจกต์แบบ SheetJS (workbook.Sheets[name][addr] = {v: ...}) ได้
#   python xlsx_shim_dump.py <input.xlsx> <output.json>
import json
import sys
import openpyxl
from openpyxl.utils import get_column_letter


def dump(in_path, out_path):
    wb = openpyxl.load_workbook(in_path, data_only=True)
    sheets_json = {}
    wb_sheets_meta = []
    for name in wb.sheetnames:
        ws = wb[name]
        wb_sheets_meta.append({"Hidden": 1 if ws.sheet_state != "visible" else 0})
        cell_map = {}
        max_r, max_c = ws.max_row, ws.max_column
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                v = ws.cell(row=r, column=c).value
                if v is not None:
                    cell_map[f"{get_column_letter(c)}{r}"] = {"v": v}
        cell_map["!ref"] = f"A1:{get_column_letter(max(max_c, 1))}{max(max_r, 1)}"
        rows_hidden = []
        for r in range(1, max_r + 1):
            rd = ws.row_dimensions.get(r)
            rows_hidden.append(1 if (rd and rd.hidden) else 0)
        cols_hidden = []
        for c in range(1, max_c + 1):
            cd = ws.column_dimensions.get(get_column_letter(c))
            cols_hidden.append(1 if (cd and cd.hidden) else 0)
        cell_map["!rows_hidden"] = rows_hidden
        cell_map["!cols_hidden"] = cols_hidden
        sheets_json[name] = cell_map
    out = {"SheetNames": wb.sheetnames, "WorkbookSheetsMeta": wb_sheets_meta, "Sheets": sheets_json}
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False)


if __name__ == "__main__":
    dump(sys.argv[1], sys.argv[2])
